#!/usr/bin/env python3
"""Fill the workspace roster from an HR or payroll export.

Usage:
  python3 match_roster.py <workspace root> <hr_export.csv> [--premium "Name, Name"]
                          [--set "email=Name|Department" ...] [--write]

The roster (config/roster.xlsx or roster.csv, bootstrapped from the spend
export) is keyed by each person's claude.ai login email. HR exports usually
aren't: they have names, often old inactive rows, repeated first names, and
contractors listed as "Allen (Fractional PM)". This matches each roster email
to one HR row and reports, for a person to decide, everything it can't match
with confidence. It never guesses between two candidates.

Matching, in order:
  1. an email column in the HR export, when it has one;
  2. the email's name part against each HR name: first.last, firstlast, flast,
     first, firstl;
  3. for first.last emails with no full-name match, a single person listed by
     first name only (contractor rows);
  then, among several candidates, the one active row if exactly one is active.

Columns are found by name (Employee/Name, Email, Department/Team,
Active/Status); rows with no department are treated as section headings.

Without --write it only reports. --premium marks those people's seats premium
(everyone else standard). --set resolves what it reported: an email, then the
HR name, then optionally a department after "|" for people HR doesn't list.
"""
from __future__ import annotations

import argparse, csv, os, re, sys, unicodedata
from typing import Iterable, Optional, Union

NAME_COLUMNS = ("employee", "name", "full name", "employee name")
EMAIL_COLUMNS = ("email", "work email", "email address")
DEPT_COLUMNS = ("department", "dept", "team")
ACTIVE_COLUMNS = ("active", "active t/f", "status", "employment status")
INACTIVE_WORDS = ("false", "inactive", "terminated", "no", "0")


def plain(text: str) -> str:
    """Lowercase ASCII letters only, so "O'Brien" and "Óbrien" compare equal."""
    ascii_text = unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z]", "", ascii_text.lower())


def pick_column(header: Iterable[str], candidates: tuple[str, ...]) -> Optional[str]:
    lowered = {h.strip().lower(): h for h in header}
    return next((lowered[c] for c in candidates if c in lowered), None)


def read_hr(path: str) -> list[dict]:
    rows = list(csv.DictReader(open(path, newline="", encoding="utf-8-sig")))
    if not rows:
        return []
    header = rows[0].keys()
    name_col, email_col = pick_column(header, NAME_COLUMNS), pick_column(header, EMAIL_COLUMNS)
    dept_col, active_col = pick_column(header, DEPT_COLUMNS), pick_column(header, ACTIVE_COLUMNS)
    if not name_col or not dept_col:
        raise SystemExit(f"Couldn't find a name and a department column in {path} (found: {', '.join(header)}).")
    people = []
    for row in rows:
        dept = (row.get(dept_col) or "").strip()
        name = (row.get(name_col) or "").strip()
        if not dept or not name:
            continue
        bare = re.sub(r"\(.*?\)", "", name).split()
        active_value = (row.get(active_col) or "").strip().lower() if active_col else ""
        people.append({
            "name": name,
            "display": " ".join(bare) if bare else name,
            "first": plain(bare[0]) if bare else "",
            "last": plain(bare[-1]) if len(bare) > 1 else "",
            "email": (row.get(email_col) or "").strip().lower() if email_col else "",
            "department": dept,
            "active": active_value not in INACTIVE_WORDS,
        })
    return people


def name_keys(person: dict) -> set[str]:
    first, last = person["first"], person["last"]
    if not last:
        return {first}
    return {f"{first}.{last}", first + last, first[0] + last, first, first + last[0]}


def prefer_active(candidates: list[dict]) -> list[dict]:
    active = [c for c in candidates if c["active"]]
    return active if len(active) == 1 else candidates


def match(email: str, people: list[dict]) -> tuple[Optional[dict], list[dict]]:
    """(person or None, candidate list) for one roster email."""
    by_email = [p for p in people if p["email"] and p["email"] == email]
    if by_email:
        found = prefer_active(by_email)
        return (found[0] if len(found) == 1 else None), found
    local = email.split("@")[0].lower()
    found = prefer_active([p for p in people if local in name_keys(p)])
    if not found and "." in local:
        first = plain(local.split(".")[0])
        found = prefer_active([p for p in people if not p["last"] and p["first"] == first])
    return (found[0] if len(found) == 1 else None), found


def load_roster(root: str) -> tuple[str, Union[str, tuple], list[list]]:
    xlsx, csv_path = os.path.join(root, "config", "roster.xlsx"), os.path.join(root, "config", "roster.csv")
    if os.path.exists(xlsx):
        import openpyxl
        book = openpyxl.load_workbook(xlsx)
        sheet = book["Roster"] if "Roster" in book.sheetnames else book.worksheets[0]
        rows = [[cell.value for cell in row[:4]] for row in sheet.iter_rows(min_row=2)]
        return "xlsx", (xlsx, book, sheet), [r for r in rows if r[0] and "@" in str(r[0])]
    if os.path.exists(csv_path):
        rows = [[r.get("email"), r.get("name"), r.get("department"), r.get("seat_type")]
                for r in csv.DictReader(open(csv_path))]
        return "csv", csv_path, [r for r in rows if r[0] and "@" in str(r[0])]
    raise SystemExit("No config/roster.xlsx or roster.csv yet — run bootstrap_roster.py on a spend export first.")


def save_roster(handle: Union[str, tuple], filled: dict[str, tuple[str, str, str]]) -> None:
    """handle is the roster.csv path, or (roster.xlsx path, workbook, sheet)."""
    if isinstance(handle, tuple):
        path, book, sheet = handle
        for row in sheet.iter_rows(min_row=2):
            email = str(row[0].value or "").lower()
            if email in filled:
                row[1].value, row[2].value, row[3].value = filled[email]
        book.save(path)
    else:
        with open(handle) as fh:
            rows = list(csv.DictReader(fh))
        for r in rows:
            if r["email"].lower() in filled:
                r["name"], r["department"], r["seat_type"] = filled[r["email"].lower()]
        with open(handle, "w", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=["email", "name", "department", "seat_type"])
            writer.writeheader()
            writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("root")
    parser.add_argument("hr_export")
    parser.add_argument("--premium", default="", help="comma-separated names (as in the HR export) with premium seats")
    parser.add_argument("--set", action="append", default=[], metavar="EMAIL=NAME[|DEPARTMENT]")
    parser.add_argument("--write", action="store_true")
    args = parser.parse_args()

    people = read_hr(args.hr_export)
    premium = {plain(n) for n in args.premium.split(",") if n.strip()}
    overrides: dict[str, tuple[str, str]] = {}
    for item in args.set:
        email, _, rest = item.partition("=")
        name, _, dept = rest.partition("|")
        overrides[email.strip().lower()] = (name.strip(), dept.strip())

    _, handle, roster = load_roster(args.root)
    filled: dict[str, tuple[str, str, str]] = {}
    ambiguous: list[tuple[str, list[str]]] = []
    unmatched: list[str] = []
    for email, *_ in roster:
        email = str(email).lower()
        if email in overrides:
            name, dept = overrides[email]
            same_name = [p for p in people if plain(p["name"]) == plain(name)]
            active = [p for p in same_name if p["active"]]
            listed: Optional[dict] = active[0] if active else (same_name[0] if same_name else None)
            if listed is None and not dept:
                unmatched.append(email)
                continue
            person = {"display": listed["display"] if listed else name,
                      "name": listed["name"] if listed else name,
                      "department": dept if dept else (listed["department"] if listed else "")}
        else:
            matched, candidates = match(email, people)
            if matched is None:
                if candidates:
                    ambiguous.append((email, [c["name"] for c in candidates]))
                else:
                    unmatched.append(email)
                continue
            person = matched
        seat = "premium" if plain(person["name"]) in premium or plain(person["display"]) in premium else "standard"
        filled[email] = (person["display"], person["department"], seat)

    found_premium = {plain(n) for n, _, s in filled.values() if s == "premium"}
    missing_premium = sorted(n.strip() for n in args.premium.split(",") if n.strip() and plain(n) not in found_premium)
    print(f"Matched {len(filled)} of {len(roster)} roster emails to {args.hr_export}.")
    print(f"Premium seats: {sum(1 for v in filled.values() if v[2] == 'premium')}"
          + (f" (not found: {', '.join(missing_premium)})" if missing_premium else ""))
    for email, names in ambiguous:
        print(f"  ? {email} could be: {' / '.join(names)}  →  --set \"{email}=<name>\"")
    for email in unmatched:
        print(f"  ✗ {email} isn't in the HR export  →  --set \"{email}=<Full Name>|<Department>\"")
    if args.write:
        save_roster(handle, filled)
        print(f"Wrote {len(filled)} people to the roster. Anyone listed above is unchanged.")
    else:
        print("Nothing written (report only). Add --write to fill the roster.")
    return 0 if not (ambiguous or unmatched or missing_premium) else 2


if __name__ == "__main__":
    sys.exit(main())
