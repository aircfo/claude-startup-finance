#!/usr/bin/env python3
"""Check an AI-spend workspace and say exactly what to do next.

Usage: python3 doctor.py <workspace root> [--offline] [--json] [--fix-config]

Runs every setup check in order and prints, for anything not ready, what is
wrong, how to fix it, and who can fix it (the operator, a claude.ai Owner, the
Console admin, or the Ramp admin). The first unfinished check is the next step,
so running it again after any interruption is how setup resumes: the state is
read from the workspace itself, never from a separate progress file that could
drift from it.

--offline   skip the checks that call Anthropic or open the export browser
--json      machine-readable output (Claude reads this during setup)
--fix-config add settings that newer versions of the skill introduced to
            config/pipeline.json, with their template defaults; nothing existing
            is changed

Each run appends one line to logs/setup.jsonl (time + each check's status), so
time from first run to first fully green run is measured without anyone timing
it. The log holds statuses only — no keys, names or spend.

Read-only except for --fix-config, the log line, and moving a key someone pasted
into PASTE_ADMIN_KEY_HERE.txt into .secrets/ (secure_key.py), so pasting the key
is all it takes. See references/setup-pitfalls.md for the failures these checks
come from.
"""
from __future__ import annotations

import argparse, csv, datetime, json, os, re, shlex, shutil, subprocess, sys, urllib.error, urllib.request
from typing import Callable, Optional

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
TEMPLATE = os.path.join(SKILL, "config.example", "pipeline.json")
TEMPLATE_SEAT_RATES = {"standard": 30, "premium": 150}
UUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)

q = shlex.quote  # every path in a suggested command, so folders with spaces still work

OPERATOR = "Operator"
OWNER = "claude.ai Owner"
CONSOLE_ADMIN = "Console admin"
RAMP_ADMIN = "Ramp admin"


class Check:
    def __init__(self, check_id: str, title: str, status: str, detail: str = "", fix: str = "", role: str = OPERATOR) -> None:
        self.id, self.title, self.status = check_id, title, status
        self.detail, self.fix, self.role = detail, fix, role

    def as_dict(self) -> dict:
        return {"id": self.id, "title": self.title, "status": self.status,
                "detail": self.detail, "fix": self.fix, "role": self.role}


def read_json(path: str) -> dict:
    """The file's JSON object, or {} when it is absent or unreadable."""
    try:
        with open(path) as fh:
            data = json.load(fh)
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def read_rows(path: str) -> list[dict]:
    try:
        return list(csv.DictReader(open(path)))
    except OSError:
        return []


def admin_key(root: str) -> str:
    key = os.environ.get("ANTHROPIC_ADMIN_API_KEY", "").strip()
    path = os.path.join(root, ".secrets", "anthropic_admin_key")
    if not key and os.path.exists(path):
        key = open(path).read().strip()
    return key


def anthropic_get(path: str, key: str) -> dict:
    req = urllib.request.Request(f"https://api.anthropic.com{path}",
                                 headers={"x-api-key": key, "anthropic-version": "2023-06-01"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        return json.load(resp)


# ---- checks, in setup order ----

def check_tools(ctx: dict) -> Check:
    missing = []
    if sys.version_info < (3, 9):
        missing.append("Python 3.9 or newer")
    if ctx["team"] and not shutil.which("uv"):
        missing.append("uv (https://docs.astral.sh/uv/)")
    if ctx["team"] and sys.platform == "darwin" and not os.path.isdir("/Applications/Google Chrome.app"):
        missing.append("Google Chrome")
    if missing:
        return Check("tools", "Tools installed", "fail", "Missing: " + ", ".join(missing),
                     "Install them, then run the doctor again.")
    return Check("tools", "Tools installed", "ok")


def check_workspace(ctx: dict) -> Check:
    missing = [d for d in ("config", "raw", ".secrets") if not os.path.isdir(os.path.join(ctx["root"], d))]
    if missing:
        return Check("workspace", "Workspace folder", "fail", f"Not set up yet (missing {', '.join(missing)}).",
                     f"python3 {q(os.path.join(HERE, 'scaffold.py'))} {q(ctx['root'])}")
    return Check("workspace", "Workspace folder", "ok", ctx["root"])


def check_config_complete(ctx: dict) -> Check:
    template = read_json(TEMPLATE)
    missing = [k for k in template if k not in ctx["cfg"] and not k.endswith("_note")]
    if missing:
        return Check("config_complete", "Settings up to date", "warn",
                     f"Settings added in newer versions are missing: {', '.join(missing)}.",
                     f"python3 {q(os.path.join(HERE, 'doctor.py'))} {q(ctx['root'])} --fix-config")
    return Check("config_complete", "Settings up to date", "ok")


def check_company(ctx: dict) -> Check:
    name = ctx["cfg"].get("company_name") or ""
    if not name or name == "Example Co":
        return Check("company", "Company name", "fail", "Still the template's \"Example Co\".",
                     "Set company_name in config/pipeline.json (or tell Claude the company name).")
    return Check("company", "Company name", "ok", name)


def check_seat_prices(ctx: dict) -> Check:
    rates = ctx["cfg"].get("seat_rates") or {}
    if not rates:
        return Check("seat_prices_confirmed", "Seat prices", "fail", "No seat_rates set.",
                     "Set seat_rates in config/pipeline.json to your plan's monthly price per seat type.")
    if rates == TEMPLATE_SEAT_RATES and not ctx["cfg"].get("seat_rates_confirmed"):
        return Check("seat_prices_confirmed", "Seat prices", "warn",
                     "Still the template's $30 standard / $150 premium.",
                     "Confirm your plan's prices (claude.ai → Settings → Billing) and set seat_rates, "
                     "or add \"seat_rates_confirmed\": true if these are right.")
    return Check("seat_prices_confirmed", "Seat prices", "ok", ", ".join(f"{k} ${v}" for k, v in rates.items()))


def check_org_id(ctx: dict) -> Check:
    if not ctx["team"]:
        return Check("org_id", "claude.ai organization id", "skip", "Enterprise plans don't need it.")
    org_id = ctx["cfg"].get("org_id") or ""
    if not UUID_RE.match(org_id):
        return Check("org_id", "claude.ai organization id", "fail",
                     "Needed for the daily export." if not org_id else f"{org_id!r} doesn't look like an organization id.",
                     "In claude.ai open Settings → Organization; copy the id from the page or the address bar "
                     "(…/organizations/<id>/…) into org_id in config/pipeline.json.", OWNER)
    return Check("org_id", "claude.ai organization id", "ok")


def check_admin_key(ctx: dict) -> Check:
    if not ctx["key"]:
        return Check("admin_key", "Admin API key", "warn",
                     "None yet. Optional: it adds Console API cost (unless Ramp supplies it) and the Claude Code report.",
                     f"In platform.claude.com → Settings → Admin keys, create a key, then paste it into "
                     f"{os.path.join(ctx['root'], 'PASTE_ADMIN_KEY_HERE.txt')} and save. Never paste it in chat.",
                     CONSOLE_ADMIN)
    return Check("admin_key", "Admin API key", "ok", "stored in .secrets/")


def check_admin_key_org(ctx: dict) -> Check:
    if not ctx["key"] or ctx["offline"]:
        return Check("admin_key_org", "Admin key's organization", "skip")
    try:
        org = anthropic_get("/v1/organizations/me", ctx["key"])
    except urllib.error.HTTPError as err:
        return Check("admin_key_org", "Admin key's organization", "fail", f"Anthropic answered HTTP {err.code}.",
                     "The key may be revoked or not an Admin key. Create a new Admin key and drop it in again.",
                     CONSOLE_ADMIN)
    except OSError as err:
        return Check("admin_key_org", "Admin key's organization", "warn", f"Couldn't reach Anthropic ({err}).",
                     "Check the internet connection and run the doctor again.")
    ctx["key_org"] = org
    detail = f"{org.get('name')} ({org.get('id')})"
    if ctx["cfg"].get("api_cost_source") == "ramp":
        return Check("admin_key_org", "Admin key's organization", "ok", detail,
                     "Ramp must hold a key for this same organization: in Ramp's AI-spend connections the "
                     "Anthropic organization id has to match. Claude checks this during setup.", RAMP_ADMIN)
    return Check("admin_key_org", "Admin key's organization", "ok", detail)


def check_claude_code_visibility(ctx: dict) -> Check:
    if not ctx["key"] or ctx["offline"]:
        return Check("claude_code_visibility", "Claude Code report covers people", "skip")
    day = (datetime.date.today() - datetime.timedelta(days=2)).isoformat()
    try:
        rows = anthropic_get(f"/v1/organizations/usage_report/claude_code?starting_at={day}&limit=100", ctx["key"]).get("data", [])
    except (OSError, urllib.error.HTTPError) as err:
        return Check("claude_code_visibility", "Claude Code report covers people", "warn", f"Couldn't read it ({err}).")
    if not rows:
        return Check("claude_code_visibility", "Claude Code report covers people", "warn",
                     f"No Claude Code usage on {day}; can't tell yet.")
    people = sum(1 for r in rows if (r.get("actor") or {}).get("type") == "user_actor")
    if not people:
        return Check("claude_code_visibility", "Claude Code report covers people", "warn",
                     "The report shows API keys only: this key's Console org is separate from your claude.ai "
                     "Team workspace, so no Admin key can see seat users.",
                     "Nothing to fix. Per-person detail comes from the spend export; the Claude Code tab will "
                     "show API keys.")
    return Check("claude_code_visibility", "Claude Code report covers people", "ok", f"{people} people on {day}")


def check_export_login(ctx: dict) -> Check:
    if not ctx["team"]:
        return Check("export_login", "Daily export can sign in", "skip", "Enterprise uses the Analytics API.")
    if not os.path.isdir(os.path.join(ctx["root"], ".browser-profile")):
        return Check("export_login", "Daily export can sign in", "fail", "The export's browser profile isn't signed in yet.",
                     f"cd {q(ctx['root'])} && uv run --with playwright python {q(os.path.join(HERE, 'fetch_surface_b_browser.py'))} "
                     f"{q(ctx['root'])} --login   (sign in to claude.ai in the window, then close it)", OWNER)
    if ctx["offline"]:
        return Check("export_login", "Daily export can sign in", "skip", "Profile present; not tested offline.")
    try:
        result = subprocess.run(["uv", "run", "--with", "playwright", "python",
                                 os.path.join(HERE, "fetch_surface_b_browser.py"), ctx["root"], "--check"],
                                capture_output=True, text=True, timeout=180)
    except (OSError, subprocess.TimeoutExpired) as err:
        return Check("export_login", "Daily export can sign in", "warn", f"Couldn't run the check ({err}).")
    message = (result.stdout.strip().splitlines() or [""])[-1]
    if result.returncode == 0:
        return Check("export_login", "Daily export can sign in", "ok", message)
    return Check("export_login", "Daily export can sign in", "fail", message,
                 f"cd {q(ctx['root'])} && uv run --with playwright python {q(os.path.join(HERE, 'fetch_surface_b_browser.py'))} "
                 f"{q(ctx['root'])} --login", OWNER)


def check_first_export(ctx: dict) -> Check:
    periods = sorted(p for p in os.listdir(os.path.join(ctx["root"], "raw"))
                     if os.path.exists(os.path.join(ctx["root"], "raw", p, "overage_spend.csv"))) \
        if os.path.isdir(os.path.join(ctx["root"], "raw")) else []
    if not periods:
        return Check("first_export", "First spend export saved", "fail", "No export in raw/ yet.",
                     f"cd {q(ctx['root'])} && uv run --with playwright python {q(os.path.join(HERE, 'fetch_surface_b_browser.py'))} {q(ctx['root'])}")
    ctx["latest_export"] = periods[-1]
    return Check("first_export", "First spend export saved", "ok", f"latest: {periods[-1]}")


def roster_people(root: str) -> list[dict]:
    """The roster the operator edits (config/roster.xlsx or .csv) when there is one;
    otherwise the config CSVs the pipeline builds from it. Reading the roster
    itself means a filled-in roster turns these checks green straight away,
    before the next pipeline run regenerates org_chart.csv and seats.csv."""
    cfg_dir = os.path.join(root, "config")
    xlsx, csv_path = os.path.join(cfg_dir, "roster.xlsx"), os.path.join(cfg_dir, "roster.csv")
    rows: list[dict] = []
    if os.path.exists(xlsx):
        try:
            import openpyxl
        except ImportError:
            openpyxl = None
        if openpyxl is not None:
            book = openpyxl.load_workbook(xlsx, read_only=True)
            sheet = book["Roster"] if "Roster" in book.sheetnames else book.worksheets[0]
            values = list(sheet.iter_rows(values_only=True))
            header = [str(h or "").strip().lower() for h in (values[0] if values else [])]
            rows = [dict(zip(header, (str(v) if v is not None else "" for v in row))) for row in values[1:]]
    elif os.path.exists(csv_path):
        rows = read_rows(csv_path)
    if not rows:
        seat_type = {r.get("email", "").lower(): r.get("seat_type", "") for r in read_rows(os.path.join(cfg_dir, "seats.csv"))}
        rows = [{**r, "seat_type": seat_type.get(r.get("email", "").lower(), "")}
                for r in read_rows(os.path.join(cfg_dir, "org_chart.csv"))]
    return [r for r in rows if "@" in (r.get("email") or "")]


def check_roster_real(ctx: dict) -> Check:
    people = roster_people(ctx["root"])
    samples = [r for r in people if r["email"].lower().endswith("@example.com")]
    if samples:
        return Check("roster_real", "Roster is your people", "fail",
                     f"{len(samples)} sample @example.com people from the template are still in config/.",
                     f"python3 {q(os.path.join(HERE, 'bootstrap_roster.py'))} {q(ctx['root'])} <YYYY-MM> to start a "
                     "real roster from your export (it replaces the samples when applied).")
    if not people:
        return Check("roster_real", "Roster is your people", "fail", "No roster yet.",
                     f"python3 {q(os.path.join(HERE, 'bootstrap_roster.py'))} {q(ctx['root'])} <YYYY-MM>")
    return Check("roster_real", "Roster is your people", "ok", f"{len(people)} people")


def check_roster_complete(ctx: dict) -> Check:
    people = roster_people(ctx["root"])
    if not people:
        return Check("roster_complete", "Roster filled in", "skip")
    no_dept = [r["email"] for r in people if not (r.get("department") or "").strip()]
    premium = sum(1 for r in people if (r.get("seat_type") or "").strip().lower() == "premium")
    if no_dept:
        return Check("roster_complete", "Roster filled in", "fail", f"{len(no_dept)} people have no department.",
                     f"Fill config/roster.xlsx, or match your HR export: python3 "
                     f"{q(os.path.join(HERE, 'match_roster.py'))} {q(ctx['root'])} <hr_export.csv>")
    if premium == 0 and "premium" in (ctx["cfg"].get("seat_rates") or {}):
        return Check("roster_complete", "Roster filled in", "warn", "No premium seats marked; everyone counts as standard.",
                     "Mark premium seats in config/roster.xlsx (or pass --premium to match_roster.py).")
    return Check("roster_complete", "Roster filled in", "ok", f"{len(people)} people, {premium} premium")


def check_roster_departments(ctx: dict) -> Check:
    counts: dict[str, int] = {}
    for r in roster_people(ctx["root"]):
        dept = (r.get("department") or "").strip()
        if dept:
            counts[dept] = counts.get(dept, 0) + 1
    singles = sorted(d for d, n in counts.items() if n == 1)
    if singles and len(counts) > 1:
        return Check("roster_departments", "Department names consistent", "warn",
                     f"One-person departments, possibly typos: {', '.join(singles)}.",
                     "If any of these should merge with another department, edit config/roster.xlsx.")
    return Check("roster_departments", "Department names consistent", "ok", f"{len(counts)} departments")


def check_dashboard(ctx: dict) -> Check:
    url = ctx["cfg"].get("dashboard_url") or ""
    if not url:
        return Check("dashboard", "Dashboard created", "fail", "No dashboard_url in config/pipeline.json.",
                     "Ask Claude: \"create the AI spend dashboard\" (references/daily-refresh.md → One-time setup).")
    return Check("dashboard", "Dashboard created", "ok", url)


def check_mac_job(ctx: dict) -> Check:
    label = ctx["cfg"].get("mac_job_label")
    if sys.platform != "darwin":
        return Check("mac_job", "Daily job scheduled", "skip", "macOS only.")
    if not label:
        return Check("mac_job", "Daily job scheduled", "fail", "Not installed.",
                     f"python3 {q(os.path.join(HERE, 'install_mac_job.py'))} {q(ctx['root'])}  (shows everything first; add --load to install)")
    loaded = subprocess.run(["launchctl", "list", label], capture_output=True).returncode == 0
    if not loaded:
        return Check("mac_job", "Daily job scheduled", "fail", f"{label} is written but not loaded.",
                     f"python3 {q(os.path.join(HERE, 'install_mac_job.py'))} {q(ctx['root'])} --load")
    return Check("mac_job", "Daily job scheduled", "ok", label)


def check_data_fresh(ctx: dict) -> Check:
    metas = []
    raw = os.path.join(ctx["root"], "raw")
    for period in sorted(os.listdir(raw)) if os.path.isdir(raw) else []:
        meta = read_json(os.path.join(raw, period, "meta.json"))
        if meta and meta.get("end"):
            metas.append(meta)
    if not metas:
        return Check("data_fresh", "Data is current", "skip")
    latest = metas[-1]
    end = datetime.date.fromisoformat(latest["end"])
    behind = (datetime.date.today() - end).days - 1
    if latest.get("fetch_error"):
        return Check("data_fresh", "Data is current", "fail", latest["fetch_error"], "See the fix for export_login.", OWNER)
    if behind > 1:
        return Check("data_fresh", "Data is current", "warn", f"Export runs through {end}, {behind} days behind.",
                     "The daily job hasn't run recently: check logs/daily.log, or that the Mac was awake.")
    return Check("data_fresh", "Data is current", "ok", f"through {end}")


CHECKS = [check_tools, check_workspace, check_config_complete, check_company, check_seat_prices, check_org_id,
          check_admin_key, check_admin_key_org, check_claude_code_visibility, check_export_login,
          check_first_export, check_roster_real, check_roster_complete, check_roster_departments,
          check_dashboard, check_mac_job, check_data_fresh]


def fix_config(root: str) -> Optional[list[str]]:
    """Add missing template settings. None, writing nothing, when the existing file
    isn't valid JSON: treating it as empty would replace the operator's settings."""
    path = os.path.join(root, "config", "pipeline.json")
    if not os.path.exists(path):
        raise SystemExit(f"[doctor] no workspace settings yet — set the workspace up first: "
                         f"python3 {q(os.path.join(HERE, 'scaffold.py'))} {q(root)}")
    try:
        with open(path) as fh:
            cfg = json.load(fh)
    except ValueError:
        return None
    added = {k: v for k, v in read_json(TEMPLATE).items() if k not in cfg}
    if added:
        cfg.update(added)
        open(path, "w").write(json.dumps(cfg, indent=2, ensure_ascii=False) + "\n")
    return sorted(k for k in added if not k.endswith("_note"))


def log_run(root: str, checks: list[Check]) -> None:
    logs = os.path.join(root, "logs")
    if not os.path.isdir(root):
        return
    os.makedirs(logs, exist_ok=True)
    line = {"at": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
            "checks": {c.id: c.status for c in checks}}
    with open(os.path.join(logs, "setup.jsonl"), "a") as fh:
        fh.write(json.dumps(line) + "\n")


def setup_minutes(root: str) -> Optional[int]:
    """Minutes from the first doctor run to the first run with no failures, or None."""
    try:
        runs = [json.loads(line) for line in open(os.path.join(root, "logs", "setup.jsonl"))]
    except (OSError, ValueError):
        return None
    green = next((r for r in runs if "fail" not in r["checks"].values()), None)
    if not runs or not green:
        return None
    start = datetime.datetime.fromisoformat(runs[0]["at"])
    return round((datetime.datetime.fromisoformat(green["at"]) - start).total_seconds() / 60)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("root")
    parser.add_argument("--offline", action="store_true")
    parser.add_argument("--json", action="store_true")
    parser.add_argument("--fix-config", action="store_true")
    args = parser.parse_args()
    root = os.path.abspath(os.path.expanduser(args.root))
    if args.fix_config:
        added = fix_config(root)
        if added is None:
            print("[doctor] config/pipeline.json isn't valid JSON (often a trailing comma); fix it first — "
                  "nothing was changed.")
            return 1
        print(f"[doctor] added settings: {', '.join(added)}" if added else "[doctor] settings already up to date")
    moved = subprocess.run([sys.executable, os.path.join(HERE, "secure_key.py"), root], capture_output=True, text=True)
    if moved.stdout.strip() and not args.json:
        print(moved.stdout.strip())  # a key pasted into PASTE_ADMIN_KEY_HERE.txt is now in .secrets/
    cfg = read_json(os.path.join(root, "config", "pipeline.json"))
    ctx = {"root": root, "cfg": cfg, "offline": args.offline, "key": admin_key(root),
           "team": (cfg.get("plan_tier") or "team").lower() == "team"}
    checks = [check(ctx) for check in CHECKS]
    log_run(root, checks)
    # Warnings never block "All set": some have nothing to fix (a report that can't see
    # seat users) or are optional. They are listed for a look, not made the next step.
    next_step = next((c for c in checks if c.status == "fail"), None)
    warnings = [c for c in checks if c.status == "warn"]
    minutes = setup_minutes(root)

    if args.json:
        print(json.dumps({"checks": [c.as_dict() for c in checks],
                          "next": next_step.as_dict() if next_step else None,
                          "warnings": [c.as_dict() for c in warnings], "setup_minutes": minutes}, indent=2))
        return 1 if any(c.status == "fail" for c in checks) else 0

    marks = {"ok": "✓", "warn": "!", "fail": "✗", "skip": "·"}
    for c in checks:
        print(f" {marks[c.status]} {c.title}" + (f" — {c.detail}" if c.detail else ""))
    waiting = sorted({c.role for c in checks if c.status == "fail" and c.role != OPERATOR})
    if next_step:
        print(f"\nNext: {next_step.title} — {next_step.detail}\n  Fix ({next_step.role}): {next_step.fix}")
    else:
        print("\nAll set." + (f" Setup took {minutes} minutes." if minutes is not None else ""))
    if warnings:
        print("\nWorth a look (not blocking): " + "; ".join(f"{c.title} — {c.fix or c.detail}" for c in warnings))
    if waiting:
        print(f"\nWaiting on someone else: {', '.join(waiting)}. Their steps are marked above.")
    return 1 if any(c.status == "fail" for c in checks) else 0


if __name__ == "__main__":
    sys.exit(main())
