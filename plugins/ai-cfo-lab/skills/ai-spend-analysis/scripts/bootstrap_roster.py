#!/usr/bin/env python3
"""Bootstrap a fill-in roster from a Surface B overage export.

Usage: bootstrap_roster.py <workspace root> <YYYY-MM>

Reads raw/<period>/overage_spend.csv, pulls the distinct people who appear in
your real usage data, and writes a fill-in roster at config/roster.xlsx (or
config/roster.csv when openpyxl is not installed). Emails and a best-effort name
guess are pre-filled; you fill in department and mark which seats are premium —
either by editing the sheet, or by pasting a roster to Claude and letting it fill
the sheet for you. apply_roster.py then turns the roster into the config CSVs the
pipeline reads.

Never overwrites an existing roster, so a filled-in one is always safe. Read-only
on the export.
"""
import csv, os, re, sys

ROSTER_COLS = ["email", "name", "department", "seat_type"]
SEAT_TYPES = ["standard", "premium"]
INSTRUCTIONS = [
    "How to fill in this roster",
    "",
    "1. One row per person who pays for a Claude seat. Emails are pre-filled from",
    "   your real spend export — add anyone who has a seat but no usage this month.",
    "2. name: correct the guess if it is wrong (it is derived from the email).",
    "3. department: type the team the person belongs to (e.g. Engineering, Sales).",
    "   Blank departments show up as '(unmapped)' on the dashboard.",
    "4. seat_type: choose 'standard' or 'premium' (Roster sheet has a dropdown).",
    "   The monthly dollar figure fills itself from seat_rates in pipeline.json —",
    "   there is no dollar column to edit here.",
    "",
    "Prefer not to type? Paste your roster into the Claude chat and ask it to fill",
    "this sheet in for you. When you are done, re-run the analysis and the",
    "dashboard rebuilds with full person / department breakdowns.",
]


def guess_name(email):
    """Best-effort display name from an email local part (user edits if wrong)."""
    local = email.split("@", 1)[0]
    parts = [p for p in re.split(r"[._\-]+", local) if p]
    return " ".join(p.capitalize() for p in parts)


def distinct_people(export_path):
    """Distinct valid emails from the overage export, sorted, with a name guess."""
    seen = {}
    with open(export_path, newline="") as fh:
        for row in csv.DictReader(fh):
            email = (row.get("user_email") or "").strip().lower()
            if "@" in email and email not in seen:
                seen[email] = {"email": email, "name": guess_name(email),
                               "department": "", "seat_type": "standard"}
    return [seen[e] for e in sorted(seen)]


def write_csv(path, people):
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=ROSTER_COLS, lineterminator="\n")
        w.writeheader()
        w.writerows(people)


def write_xlsx(path, people):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roster"
    header_font = Font(bold=True)
    for col, name in enumerate(ROSTER_COLS, 1):
        ws.cell(1, col, name).font = header_font
    for i, person in enumerate(people, 2):
        for col, name in enumerate(ROSTER_COLS, 1):
            ws.cell(i, col, person[name])
    # Dropdown on seat_type keeps the values valid without the user memorizing them.
    dv = DataValidation(type="list", formula1=f'"{",".join(SEAT_TYPES)}"', allow_blank=False)
    ws.add_data_validation(dv)
    last = len(people) + 1
    dv.add(f"D2:D{max(last, 2)}")
    highlight = PatternFill("solid", fgColor="FFF3CD")
    for row in range(2, last + 1):
        ws.cell(row, 3).fill = highlight  # department — the column the user must fill
    ws.freeze_panes = "A2"
    for col, name in enumerate(ROSTER_COLS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = max(16, len(name) + 22)

    readme = wb.create_sheet("How to fill this in")
    for i, line in enumerate(INSTRUCTIONS, 1):
        cell = readme.cell(i, 1, line)
        if i == 1:
            cell.font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 82
    wb.save(path)


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    root, period = sys.argv[1], sys.argv[2]
    export_path = os.path.join(root, "raw", period, "overage_spend.csv")
    if not os.path.exists(export_path) or os.path.getsize(export_path) == 0:
        print(f"[roster] no export at {export_path} — pull the spend export first, then re-run", file=sys.stderr)
        return 1

    cfg = os.path.join(root, "config")
    os.makedirs(cfg, exist_ok=True)
    xlsx_path, csv_path = os.path.join(cfg, "roster.xlsx"), os.path.join(cfg, "roster.csv")
    existing = next((p for p in (xlsx_path, csv_path) if os.path.exists(p)), None)
    if existing:
        print(f"[roster] {os.path.relpath(existing, root)} already exists — leaving it untouched. "
              f"Edit it, or delete it to re-bootstrap from the export.")
        return 0

    people = distinct_people(export_path)
    if not people:
        print(f"[roster] no valid emails found in {os.path.relpath(export_path, root)}", file=sys.stderr)
        return 1

    try:
        import openpyxl  # noqa: F401
        write_xlsx(xlsx_path, people)
        target = xlsx_path
    except ImportError:
        write_csv(csv_path, people)
        target = csv_path

    rel = os.path.relpath(target, root)
    print(f"[roster] wrote {rel} with {len(people)} people from the {period} export.")
    print("  → fill in department + seat_type (or paste your roster to Claude to fill it), "
          "then re-run the analysis. apply_roster.py turns it into org_chart.csv + seats.csv.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
